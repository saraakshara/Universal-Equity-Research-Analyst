from __future__ import annotations
from io import BytesIO
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet


def excel_bytes(payload):
    bio=BytesIO(); wb=Workbook(); ws=wb.active; ws.title="Executive Summary"
    rows=payload["summary"]
    for r,row in enumerate(rows,1):
        for c,v in enumerate(row,1): ws.cell(r,c,v)
    for sname, df in payload.get("sheets", {}).items():
        ws=wb.create_sheet(sname[:31])
        if isinstance(df,pd.DataFrame):
            data=df.reset_index()
            for c,col in enumerate(data.columns,1): ws.cell(1,c,str(col)).font=Font(bold=True)
            for r,row in enumerate(data.itertuples(index=False),2):
                for c,v in enumerate(row,1): ws.cell(r,c,v if v==v else None)
    wb.save(bio); return bio.getvalue()


def pdf_bytes(payload):
    bio=BytesIO(); doc=SimpleDocTemplate(bio,pagesize=A4,rightMargin=36,leftMargin=36,topMargin=36,bottomMargin=36)
    styles=getSampleStyleSheet(); story=[]
    story.append(Paragraph(payload["title"],styles["Title"])); story.append(Spacer(1,10))
    for heading, body in payload.get("sections",[]):
        story.append(Paragraph(heading,styles["Heading2"]))
        for line in body:
            story.append(Paragraph(str(line).replace("&","&amp;"),styles["BodyText"])); story.append(Spacer(1,4))
    doc.build(story); return bio.getvalue()
